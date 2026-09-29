"""test_network_discovery.py — Network Discovery Framework v1."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import target_input as ti
import network_discovery as nd
import asset_inventory as ainv
import asset_graph as ag
import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


# ── 1. Target Type Detection / Normalization ────────────────────────────────
def test_target_type_detection():
    assert ti.detect_target_type("https://example.com/path") == ti.T_URL
    assert ti.detect_target_type("example.com") == ti.T_DOMAIN
    assert ti.detect_target_type("192.168.0.10") == ti.T_IP
    assert ti.detect_target_type("192.168.0.10,192.168.0.11") == ti.T_IP_LIST
    assert ti.detect_target_type("192.168.0.0/24") == ti.T_CIDR


def test_cidr_normalization():
    out = ti.normalize_targets("192.168.0.0/24")
    assert out["target_type"] == ti.T_CIDR
    item = out["normalized_targets"][0]
    assert item["type"] == ti.T_CIDR
    assert item["host_count"] == 254          # /24 → 254 usable
    assert item["is_private"] is True
    assert out["scan_scope_summary"]["estimated_hosts"] == 254


def test_private_public_split():
    out = ti.normalize_targets("8.8.8.8,10.0.0.5")
    s = out["scan_scope_summary"]
    assert s["public_hosts"] == 1 and s["private_hosts"] == 1


# ── 2. Scope Guard / Max Host / 큰 대역 차단 ─────────────────────────────────
def test_scope_guard_allows_slash24():
    norm = ti.normalize_targets("192.168.1.0/24")
    g = nd.scope_guard(norm)
    assert g["within_limit"] is True
    assert g["total_hosts"] == 254


def test_scope_guard_blocks_large_range():
    norm = ti.normalize_targets("10.0.0.0/8")     # 거대 대역
    g = nd.scope_guard(norm)
    assert g["within_limit"] is False
    assert len(g["blocked"]) == 1
    assert "큰 대역" in g["blocked_reasons"][0] or "차단" in g["blocked_reasons"][0]


def test_scope_guard_max_hosts(monkeypatch):
    monkeypatch.setenv("NETWORK_MAX_HOSTS", "100")
    norm = ti.normalize_targets("192.168.0.0/24")  # 254 > 100
    g = nd.scope_guard(norm, nd.network_policy())
    assert g["within_limit"] is False


def test_large_range_allowed_with_explicit_approval(monkeypatch):
    monkeypatch.setenv("NETWORK_ALLOW_LARGE", "true")
    norm = ti.normalize_targets("10.0.0.0/16")
    g = nd.scope_guard(norm, nd.network_policy())
    assert g["within_limit"] is True


# ── Live Host Discovery (주입형, 네트워크 없음) ──────────────────────────────
def test_expand_and_discover_live_hosts():
    norm = ti.normalize_targets("192.168.0.0/30")  # 2 usable hosts
    g = nd.scope_guard(norm)
    ips = nd.expand_targets(g["allowed"], 256)
    assert len(ips) == 2
    # 첫 IP만 살아있다고 mock
    def fake_ping(ip, timeout=2.0):
        return ip == ips[0]
    disc = nd.discover_live_hosts(ips, ping_fn=fake_ping)
    assert disc["live_hosts"] == [ips[0]]
    assert disc["unreachable_hosts"] == [ips[1]]
    assert disc["discovery_summary"]["live"] == 1


def test_run_discovery_end_to_end():
    out = nd.run_discovery("192.168.5.0/30", ping_fn=lambda ip, t=2.0: True)
    assert out["target_type"] == ti.T_CIDR
    assert out["network_discovery"]["discovery_summary"]["live"] == 2


# ── 4~5. Asset Inventory / Classification ────────────────────────────────────
def _host_results():
    return [
        {"host": "10.0.0.5", "ip": "10.0.0.5", "open_ports": [3306, 22],
         "services": [{"port": 3306, "service": "mysql", "version": "5.7"},
                      {"port": 22, "service": "ssh"}]},
        {"host": "8.8.4.4", "ip": "8.8.4.4", "open_ports": [80, 443],
         "services": [{"port": 80, "service": "http"}, {"port": 443, "service": "https"}]},
        {"host": "10.0.0.9", "ip": "10.0.0.9", "open_ports": [6379],
         "services": [{"port": 6379, "service": "redis", "banner": "redis noauth"}]},
    ]


def test_asset_inventory_and_classification():
    inv = ainv.build_inventory(_host_results())
    by_ip = {a["ip"]: a for a in inv["assets"]}
    assert by_ip["10.0.0.5"]["asset_type"] == ainv.A_DB
    assert by_ip["8.8.4.4"]["asset_type"] == ainv.A_WEB
    assert by_ip["10.0.0.9"]["asset_type"] == ainv.A_DB   # redis = data exposure → DB군
    # risk tags
    assert "exposed_database" in by_ip["10.0.0.5"]["risk_tags"]
    assert "exposed_auth_service" in by_ip["10.0.0.5"]["risk_tags"]
    assert "internet_exposed" in by_ip["8.8.4.4"]["risk_tags"]
    assert "internet_exposed" not in by_ip["10.0.0.5"]["risk_tags"]
    # web endpoints
    assert "https://8.8.4.4:443" in by_ip["8.8.4.4"]["web_endpoints"]


def test_asset_inventory_summary():
    inv = ainv.build_inventory(_host_results())
    s = inv["summary"]
    assert s["total_assets"] == 3
    assert s["web_servers"] == 1
    assert s["database_servers"] == 2
    assert s["web_candidates"] >= 1


# ── 6. Asset Graph ───────────────────────────────────────────────────────────
def test_asset_graph_nodes_edges():
    inv = ainv.build_inventory(_host_results())
    norm = ti.normalize_targets("10.0.0.0/24")
    g = ag.build_asset_graph(inv["assets"], norm["normalized_targets"], {})
    types = {n["node_type"] for n in g["asset_nodes"]}
    assert ag.N_RANGE in types and ag.N_HOST in types and ag.N_PORT in types
    rels = {e["relation"] for e in g["asset_edges"]}
    assert ag.E_CONTAINS in rels and ag.E_EXPOSES in rels and ag.E_RUNS in rels
    assert g["asset_summary"]["hosts"] == 3


# ── 8~9. Report rendering ─────────────────────────────────────────────────────
def test_report_network_sections_render():
    inv = ainv.build_inventory(_host_results())
    a = {"asset_inventory": inv["assets"], "asset_summary": inv["summary"],
         "network_exposure_summary": {
             "input_targets": ["10.0.0.0/24"], "target_type": "cidr",
             "live_hosts": 3, "open_ports": 5, "services": 3, "web_candidates": 2,
             "auth_services": 1, "db_services": 2, "file_sharing_services": 0,
             "container_services": 0, "unknown_services": 0, "internet_exposed": 1,
             "high_priority_assets": 2, "blocked_ranges": 0}}
    doc = Document()
    report._add_v2_network_exposure(doc, a)
    report._add_v2_asset_inventory(doc, a)
    report._add_v2_dashboard(doc, a, "10.0.0.0/24")
    txt = _txt(doc)
    assert "네트워크 노출 요약" in txt
    assert "자산 인벤토리" in txt
    assert "Live Host" in txt          # dashboard KPI
    assert "10.0.0.5" in txt


def test_report_skips_without_network_data():
    doc = Document()
    report._add_v2_network_exposure(doc, {})
    report._add_v2_asset_inventory(doc, {})
    assert all(not p.text for p in doc.paragraphs)


# ── 안전 기본값 ──────────────────────────────────────────────────────────────
def test_conservative_defaults():
    p = nd.network_policy()
    assert p["max_hosts"] == 256      # /24
    assert p["block_prefix"] == 16    # /16 이상 차단
    assert p["allow_large"] is False
