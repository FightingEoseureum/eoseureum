"""AI 사용/보고서 표시 정책 회귀 테스트.

검증 항목:
- ENABLE_AI_ANALYSIS=false → enrich 후 finding 에 ai_analysis 없음(또는 used=False),
  보고서의 취약점 상세에 "AI 보안 분석가 의견" 미출력.
- ai_status.used=True + finding.ai_analysis.provider=="ollama" → 취약점 상세에
  "AI 보안 분석가 의견" 출력.
- AI 보강이 severity / confidence_score / finding 개수를 변경하지 않음(deepcopy 비교).
- Executive Summary 에 ai_status 가 1회만 표시(미사용 시 fallback_reason 포함).

실행:
  venv_linux/bin/python -m pytest tests/test_ai_report_visibility.py -q
"""
import asyncio
import copy
import io

import pytest
from docx import Document

import ai_provider
import report


# ── 헬퍼 ────────────────────────────────────────────────────────────────────────

def _base_finding(title, **extra):
    f = {
        "title": title,
        "host": "example.com",
        "port": 443,
        "service": "https",
        "severity": "MEDIUM",
        "report_severity": "Medium",
        "confidence": "CONFIRMED",
        "confidence_score": 80,
        "judgment": "취약",
        "owasp": "A03",
        "finding_type": "web",
        "recommendation": "보안 설정을 점검하십시오.",
        "evidence_detail": "근거 텍스트입니다.",
    }
    f.update(extra)
    return f


def _build_scan(findings, summary=None, **analysis_extra):
    analysis = {
        "overall_risk": "MEDIUM",
        "overall_summary": "합성 점검 결과 요약입니다.",
        "findings": findings,
        "summary": summary or {},
    }
    analysis.update(analysis_extra)
    return {
        "domain": "example.com",
        "created_at": "2026-06-18",
        "results": [
            {
                "host": "example.com",
                "ip": "192.0.2.1",
                "scan_mode": "standard",
                "open_ports": [443],
                "services": [],
            }
        ],
        "analysis": analysis,
    }


def _full_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ── 1) ENABLE_AI_ANALYSIS=false ─────────────────────────────────────────────────

def test_disabled_no_ai_analysis_on_findings(monkeypatch):
    monkeypatch.setenv("ENABLE_AI_ANALYSIS", "false")
    findings = [_base_finding("반사형 XSS 취약점")]
    analysis = {"findings": findings}

    out = asyncio.run(ai_provider.enrich_with_security_analyst("example.com", analysis))

    # finding 에 ai_analysis 가 붙지 않음.
    assert "ai_analysis" not in findings[0]
    # ai_status 는 기록되며 used=False, 사유는 ENABLE_AI_ANALYSIS_FALSE.
    st = out["ai_status"]
    assert st["enabled"] is False
    assert st["used"] is False
    assert st["fallback_reason"] == ai_provider.FALLBACK_ENABLE_FALSE


def test_disabled_report_has_no_per_finding_ai_section(monkeypatch):
    monkeypatch.delenv("SHOW_AI_SECTION_ONLY_WHEN_USED", raising=False)
    monkeypatch.delenv("SHOW_RULE_BASED_ANALYSIS", raising=False)
    findings = [_base_finding("반사형 XSS 취약점")]
    summary = {"vulnerability_count": 1, "confirmed_count": 1}
    analysis = {
        "overall_risk": "MEDIUM",
        "overall_summary": "요약",
        "findings": findings,
        "summary": summary,
        # AI 미사용 상태.
        "ai_status": {
            "enabled": False, "provider": "", "model": "", "base_url": "",
            "available": False, "used": False, "fallback": True,
            "fallback_reason": ai_provider.FALLBACK_ENABLE_FALSE,
            "last_error": "", "analyzed_findings": 0,
        },
    }
    scan = {"domain": "example.com", "created_at": "2026-06-18",
            "results": [{"host": "example.com", "ip": "192.0.2.1",
                         "scan_mode": "standard", "open_ports": [443], "services": []}],
            "analysis": analysis}

    buf = report.generate_report(scan)
    text = _full_text(Document(io.BytesIO(buf.getvalue())))

    # 취약점 상세에 AI 의견 미출력.
    assert "AI 보안 분석가 의견" not in text
    # Executive Summary 에 미사용 + 사유 1회 표시.
    assert text.count("AI 분석: 미사용") == 1
    assert "AI 분석 비활성화" in text


# ── 2) ai_status.used=True + ollama finding ─────────────────────────────────────

def test_used_report_shows_per_finding_ai_section():
    ai_block = {
        "provider": "ollama",
        "model": "llama3.2:3b",
        "used_rag": True,
        "false_positive_assessment": "오탐 가능성 낮음.",
        "business_impact": "세션 탈취 가능.",
        "attack_chain_analysis": "반사 입력을 통한 스크립트 실행.",
        "remediation_priority_reason": "사용자 영향이 크므로 우선 조치.",
        "additional_verification_steps": ["입력 인코딩 검증"],
        "report_text": "반사형 XSS 에 대한 방어적 분석.",
    }
    f = _base_finding("반사형 XSS 취약점", ai_analysis=ai_block)
    summary = {"vulnerability_count": 1, "confirmed_count": 1}
    analysis = {
        "overall_risk": "MEDIUM",
        "overall_summary": "요약",
        "findings": [f],
        "summary": summary,
        "ai_status": {
            "enabled": True, "provider": "ollama/llama3.2:3b", "model": "llama3.2:3b",
            "base_url": "http://localhost:11434", "available": True, "used": True,
            "fallback": False, "fallback_reason": "", "last_error": "",
            "analyzed_findings": 1,
        },
    }
    scan = {"domain": "example.com", "created_at": "2026-06-18",
            "results": [{"host": "example.com", "ip": "192.0.2.1",
                         "scan_mode": "standard", "open_ports": [443], "services": []}],
            "analysis": analysis}

    buf = report.generate_report(scan)
    text = _full_text(Document(io.BytesIO(buf.getvalue())))

    assert "AI 보안 분석가 의견" in text
    assert "세션 탈취 가능." in text
    # Executive Summary 에 '사용' 1회 표시.
    assert text.count("AI 분석: 사용") == 1


# ── 3) AI 보강이 finding 을 변경하지 않음 (읽기 전용) ────────────────────────────

class _FakeOllamaProvider:
    """ollama 처럼 동작하지만 고정 JSON 을 반환하는 mock provider."""
    name = "ollama/llama3.2:3b"
    base_url = "http://localhost:11434"
    model = "llama3.2:3b"

    async def is_available(self):
        return True

    async def complete(self, prompt):
        return (
            '{"false_positive_assessment": "낮음",'
            ' "business_impact": "영향 있음",'
            ' "attack_chain_analysis": "체인",'
            ' "remediation_priority_reason": "우선",'
            ' "additional_verification_steps": ["검증"],'
            ' "report_text": "요약",'
            ' "executive_summary": "경영진 요약",'
            ' "overall_risk_commentary": "코멘트",'
            ' "top_priorities": ["A"],'
            ' "attack_chain_summary": "요약",'
            ' "operation_team_actions": ["운영"],'
            ' "developer_team_actions": ["개발"],'
            ' "severity": "HIGH", "confidence_score": 1}'
        )


def test_ai_does_not_mutate_finding_fields(monkeypatch):
    monkeypatch.setenv("ENABLE_AI_ANALYSIS", "true")
    monkeypatch.setattr(ai_provider, "get_ai_provider", lambda: _FakeOllamaProvider())

    findings = [
        _base_finding("반사형 XSS 취약점", severity="MEDIUM", confidence_score=80),
        _base_finding("SQL 인젝션 취약점", severity="HIGH", confidence_score=95,
                      report_severity="High"),
    ]
    analysis = {"findings": findings}

    # AI 가 건드리면 안 되는 핵심 필드 스냅샷.
    before = copy.deepcopy(
        [{k: f.get(k) for k in
          ("severity", "confidence", "confidence_score", "judgment",
           "finding_type", "report_severity")}
         for f in findings]
    )
    before_count = len(findings)

    out = asyncio.run(ai_provider.enrich_with_security_analyst("example.com", analysis))

    after = [{k: f.get(k) for k in
              ("severity", "confidence", "confidence_score", "judgment",
               "finding_type", "report_severity")}
             for f in out["findings"]]

    assert after == before
    assert len(out["findings"]) == before_count
    # 사용된 경우 ai_analysis(provider=ollama) 부여 + 카운트.
    st = out["ai_status"]
    assert st["used"] is True
    assert st["analyzed_findings"] == before_count
    for f in out["findings"]:
        assert f["ai_analysis"]["provider"] == "ollama"


# ── 4) Executive Summary 의 ai_status 1회 표시 (미사용 + fallback_reason) ─────────

def test_exec_summary_ai_status_once_when_unused():
    analysis = {
        "overall_risk": "GOOD",
        "overall_summary": "요약",
        "findings": [],
        "summary": {"vulnerability_count": 0},
        "ai_status": {
            "enabled": True, "provider": "ollama/llama3.2:3b", "model": "llama3.2:3b",
            "base_url": "http://localhost:11434", "available": False, "used": False,
            "fallback": True, "fallback_reason": ai_provider.FALLBACK_CONN_FAILED,
            "last_error": "연결 실패", "analyzed_findings": 0,
        },
    }
    scan = {"domain": "example.com", "created_at": "2026-06-18",
            "results": [{"host": "example.com", "ip": "192.0.2.1",
                         "scan_mode": "standard", "open_ports": [443], "services": []}],
            "analysis": analysis}

    buf = report.generate_report(scan)
    text = _full_text(Document(io.BytesIO(buf.getvalue())))

    assert text.count("AI 분석: 미사용") == 1
    assert "Ollama 서버 연결 실패" in text


# ── 5) 점검 커버리지 요약 섹션 존재 ──────────────────────────────────────────────

def test_coverage_summary_section_present():
    scan = _build_scan([_base_finding("반사형 XSS 취약점")],
                       summary={"vulnerability_count": 1})
    buf = report.generate_report(scan)
    text = _full_text(Document(io.BytesIO(buf.getvalue())))
    assert "점검 커버리지 요약" in text
    assert "AI fallback reason" in text
    assert "USE_PROBE_ORCHESTRATOR" in text


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-q"]))
