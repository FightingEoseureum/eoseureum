"""test_evidence_report.py — Attack Surface 섹션 / 경영진 검증요약 / 보안 관찰 블록 렌더링."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_attack_surface_section_renders():
    analysis = {"attack_surface_plan": {
        "summary": {"total_surfaces": 20, "selected_surfaces": 8,
                    "auth_after_surfaces": 5, "critical_high_surfaces": 6,
                    "by_type": {"Object Reference": 4, "Search": 3}, "by_level": {}},
        "reasoning_trace": [
            {"input": "userId", "semantic_role": "object_reference",
             "recommended": ["idor", "access_control"], "priority": "Critical",
             "reason": "객체 식별자이며 IDOR 검증 가치 높음"},
        ],
    }}
    doc = Document()
    report._add_attack_surface_plan_section(doc, analysis)
    txt = _txt(doc)
    assert "공격 표면 분석 및 우선순위" in txt
    assert "인증 후 공격 표면 수" in txt
    assert "AI Reasoning Trace" in txt
    assert "userId" in txt and "object_reference" in txt


def test_evidence_level_executive_summary():
    analysis = {
        "evidence_levels": {"level3_proven": 2, "level2_evidence": 3, "level1_observed": 5,
                            "level0_info": 10, "priority_actions": 5},
        "attack_surface_plan": {"summary": {"total_surfaces": 20, "auth_after_surfaces": 5}},
    }
    doc = Document()
    report._add_evidence_level_summary(doc, analysis)
    txt = _txt(doc)
    assert "보안 관찰 및 검증 요약" in txt
    assert "Level 3 · 실증 확인" in txt
    assert "우선 조치 항목 수" in txt
    assert "Rule Engine" in txt


def test_observed_view_block_renders():
    f = {"title": "IDOR 실증 — 교차 계정", "confidence": "CONFIRMED_RESPONSE",
         "evidence_detail": "계정 B 응답에 계정 A 식별정보 노출",
         "business_impact": "타 사용자 정보 접근", "recommendation": "권한 검증"}
    doc = Document()
    report._add_observed_view_block(doc, f)
    txt = _txt(doc)
    assert "Level 3" in txt
    assert "관찰 내용" in txt and "확보 증거" in txt and "비즈니스 영향" in txt


def test_sections_skip_when_no_data():
    doc = Document()
    n = len(doc.paragraphs)
    report._add_attack_surface_plan_section(doc, {})
    report._add_evidence_level_summary(doc, {})
    assert len(doc.paragraphs) == n
