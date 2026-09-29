"""test_validation_readiness.py — Validation-Centric: 후보/우선순위/Backlog/점수개선/통제/요약."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import validation_readiness as vr
import attack_path_prioritizer as app
import attack_graph as ag
import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _analysis():
    return {
        "domain": "t.example.com",
        "findings": [
            {"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER", "evidence_detail": "alert"},
        ],
        "discovery_items": [
            {"title": "[참고] IDOR 가능성 — 객체 참조", "confidence": "MANUAL_REVIEW"},   # L1
            {"title": "[참고] SQL 인젝션 가능성", "confidence": "POSSIBLE"},               # L2
            {"title": "[참고] CSRF 가능성", "confidence": "MANUAL_REVIEW"},                # L1, cap L2
        ],
        "evidence_levels": {"level3_proven": 1, "level2_evidence": 1, "level1_observed": 2,
                            "level0_info": 0},
    }


# ── Validation Candidate 생성 + 우선순위 ─────────────────────────────────────
def test_validation_candidates_created():
    out = vr.build_validation_readiness(_analysis())
    fams = {c["family"] for c in out["validation_candidates"]}
    assert "idor" in fams and "sqli" in fams and "csrf" in fams
    # 이미 실증된 XSS(L3)는 후보 아님
    assert "xss" not in fams


def test_validation_priority_and_target():
    out = vr.build_validation_readiness(_analysis())
    by = {c["family"]: c for c in out["validation_candidates"]}
    # IDOR(L1) → 실증(L3) 가능, 고영향 → High
    assert by["idor"]["target_level"] == 3
    assert by["idor"]["validation_priority"] == "High"
    assert by["idor"]["promotable_to_level3"] is True
    # SQLi(L2) → L3 가능
    assert by["sqli"]["target_level"] == 3
    # CSRF(L1) → 최대 L2(수동), 실증 불가
    assert by["csrf"]["target_level"] == 2
    assert by["csrf"]["promotable_to_level3"] is False
    assert by["csrf"]["validation_priority"] == "Low"


def test_validation_backlog_summary():
    out = vr.build_validation_readiness(_analysis())
    s = out["validation_summary"]
    assert s["current_confirmed"] == 1
    # IDOR(L1→L3)·SQLi(L2→L3) 2건이 실증 가능 → potential = 1 + 2 = 3
    assert s["expected_promotions"] == 2
    assert s["potential_confirmed"] == 3
    assert s["level1_to_level2"] >= 1
    assert s["level2_to_level3"] >= 1
    assert "by_family" in s
    assert out["validation_backlog"] == out["validation_candidates"]


def test_validation_recommended_methods():
    out = vr.build_validation_readiness(_analysis())
    by = {c["family"]: c for c in out["validation_candidates"]}
    assert "A/B" in by["idor"]["recommended_validation"]
    assert "교차 계정" in by["idor"]["required_evidence"]
    assert "SQLMap" in by["sqli"]["recommended_validation"]


# ── Attack Path Evidence Weight (실증 경로 > 관찰 경로) ──────────────────────
def test_evidence_weight_confirmed_outranks_observed_admin():
    # 실증 XSS(Level3, Confirmed) vs 관찰 관리자 노출(Level1, Observed)
    a = {"domain": "t", "coverage": {},
         "findings": [{"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER",
                       "evidence_detail": "alert", "evidence_url": "http://t/s?q=1"}],
         "attack_surface_items": [{"title": "관리자 페이지 노출(401)", "confidence": "MANUAL_REVIEW"}],
         "discovery_items": []}
    g = ag.build_attack_graph(a)
    nodes = g["attack_graph"]["nodes"]
    idx = app._nodes_index(nodes)
    xss = next(p for p in g["attack_paths"] if "XSS" in p["title"] or "스크립트" in p["title"])
    admin = next((p for p in g["attack_paths"] if "관리자" in p["title"]), None)
    sx = app.score_path(xss, idx)["score"]
    assert sx >= app._W_LEVEL3   # Level3 가중 반영
    if admin:
        sa = app.score_path(admin, idx)["score"]
        assert sx > sa            # 실증 경로가 관찰 경로보다 항상 높음


def test_score_uses_path_confidence_weight():
    p = {"node_ids": ["n1"], "evidence_level": 3, "path_confidence": "Confirmed Path"}
    idx = {"n1": {"node_id": "n1", "node_type": ag.N_FINDING}}
    sc = app.score_path(p, idx)
    # Level3(+15) + Confirmed(+10) = 25 이상 → Critical
    assert sc["score"] >= 25
    assert sc["priority"] == "Critical"


# ── Security Controls 강화(실증 실패/미발견 자동 수집) ───────────────────────
def test_security_controls_collect_failures():
    a = {
        "findings": [{"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER"}],
        "discovery_items": [
            {"title": "[참고] SQL 인젝션 가능성", "confidence": "MANUAL_REVIEW"},  # 미실증
        ],
        "candidate_verification": {"idor": {"verified": 1, "promoted": 0}},
        "good_items": [], "attack_surface_items": [],
    }
    controls = report._derive_verified_controls(a)
    joined = " ".join(controls)
    assert "SQL Injection 실증 실패" in joined
    assert "교차 계정 접근 미확인" in joined


# ── Validation Opportunities 렌더링 + Executive Summary KPI ──────────────────
def test_validation_dashboard_renders():
    a = _analysis()
    out = vr.build_validation_readiness(a)
    a["validation_candidates"] = out["validation_candidates"]
    a["validation_summary"] = out["validation_summary"]
    doc = Document()
    report._add_v2_validation_opportunities(doc, a)
    txt = _txt(doc)
    assert "추가 검증 가능 항목" in txt
    assert "현재 실증(L3)" in txt
    assert "추천 검증" in txt
    assert "교차 계정" in txt


def test_validation_dashboard_skips_without_data():
    doc = Document()
    report._add_v2_validation_opportunities(doc, {})
    assert all(not p.text for p in doc.paragraphs)


def test_deterministic_no_ai():
    out = vr.build_validation_readiness(_analysis())
    assert out["validation_candidates"] and out["validation_summary"]
