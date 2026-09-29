"""test_report_v5_design.py — Report UX/UI Framework v5 (디자인 완성도)."""
import os, sys, io
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report
from docx import Document
from docx.oxml.ns import qn


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _rich():
    from tests.test_report_v4_ux import _rich as r
    return r()


def _gen():
    a = _rich()
    buf = report.generate_report({"domain": "t.example.com", "analysis": a,
                                  "created_at": "2026-07-01", "results": []})
    return Document(io.BytesIO(buf.getvalue()))


# ── 표지 레이아웃 ─────────────────────────────────────────────────────────────
def test_cover_layout():
    doc = _gen(); txt = _txt(doc)
    assert "Security Assessment Report" in txt
    assert "Web · Service · Network Security Assessment" in txt
    assert "C O N F I D E N T I A L" in txt
    assert "Prepared by" in txt and "Eoseureum" in txt
    assert "Report No" in txt and "Version" in txt


# ── 목차(점선 Leader) ────────────────────────────────────────────────────────
def test_toc_dot_leader():
    doc = _gen()
    # 점선 Leader tab stop 이 목차 문단에 존재
    found = False
    for p in doc.paragraphs:
        for ts in p.paragraph_format.tab_stops:
            if ts.leader is not None and str(ts.leader).upper().find("DOT") >= 0:
                found = True
    txt = _txt(doc)
    assert "Contents" in txt and "Executive Summary" in txt and "Recommendation" in txt
    assert found, "목차 점선 Leader tab stop 없음"


# ── Header / Footer (페이지 번호 필드) ───────────────────────────────────────
def test_header_footer_page_field():
    doc = _gen()
    sec = doc.sections[0]
    htxt = " ".join(p.text for p in sec.header.paragraphs)
    assert "CONFIDENTIAL" in htxt
    # 푸터에 PAGE 필드(instrText) 존재
    ftr_xml = sec.footer._element.xml
    assert "PAGE" in ftr_xml
    ftxt = " ".join(p.text for p in sec.footer.paragraphs)
    assert "Eoseureum · AI Security Assessment Report" in ftxt   # 보고서명(좌) + 페이지(중) + 기관(우)


# ── Finding Card: 4박스 ──────────────────────────────────────────────────────
def test_finding_card_four_boxes():
    a = _rich(); doc = Document()
    report._add_observed_view_block(doc, a["findings"][0], a)
    txt = _txt(doc)
    for box in ("관찰 내용", "확보 증거", "비즈니스 영향", "권장 조치"):
        assert f"{box}" in txt
    # 상단 배지: Severity + CVSS/CWE/OWASP 중 하나
    assert "HIGH" in txt


# ── Risk Matrix: 취약점 이름 표시 ────────────────────────────────────────────

# ── Attack Story: 도형 Flow(표) + 화살표, 엔진 용어 없음 ─────────────────────

# ── 빈 페이지 제거: 연속 빈 문단 축소 ────────────────────────────────────────
def test_no_excessive_blank_paragraphs():
    doc = _gen()
    # 연속 빈 문단(텍스트/런 없음)이 2개 초과로 이어지지 않아야 함
    run = 0; mx = 0
    for p in doc.paragraphs:
        empty = (not p.text.strip()) and (not p._p.findall(qn("w:r")))
        run = run + 1 if empty else 0
        mx = max(mx, run)
    assert mx <= 2, f"연속 빈 문단 {mx}개(빈 페이지 위험)"


# ── 여백 최적화 ──────────────────────────────────────────────────────────────
def test_margins_optimized():
    doc = _gen()
    sec = doc.sections[0]
    # v5: 여백 축소(<= 2.0cm 상하)
    from docx.shared import Cm
    assert sec.top_margin <= Cm(2.0) and sec.left_margin <= Cm(2.3)


# ── 색상 시스템(5단계) / 정규화 ───────────────────────────────────────────────
def test_sev5_color_system():
    assert set(report.SEV5_HEX.keys()) == {"Critical", "High", "Medium", "Low", "Info"}
    assert report._sev5("높음") == "High" and report._sev5("HIGH") == "High"
    assert report._sev5("긴급") == "Critical"
    # 같은 severity → 항상 같은 색
    assert report.SEV5_HEX["High"] == report.SEV5_HEX[report._sev5("높음")]


# ── 긴 문자열 접기 ───────────────────────────────────────────────────────────
def test_long_string_folded():
    long = "A" * 500
    out = report._fold_long(long, 100)
    assert len(out) < 200 and "생략" in out


# ── 전체 회귀 ─────────────────────────────────────────────────────────────────
def test_full_report_v5_intact(monkeypatch):
    monkeypatch.setenv("REPORT_SHOW_ENGINE_INTERNAL", "true")   # 풀 리포트(엔진내부 포함) 검증
    doc = _gen(); txt = _txt(doc)
    for sect in ("Security Assessment Report", "Contents", "Executive Summary",
                 "보안 진단 대시보드", "점검 목적", "취약점 상세", "발견된 취약점",
                 "개선 로드맵", "부록", "Developer Appendix"):
        assert sect in txt, f"섹션 누락: {sect}"
