"""test_report_v3_ux.py — Report UX/UI Framework v3 (표현 개편, 데이터 유지)."""
import os, sys, io
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _rich():
    return {
        "domain": "t.example.com", "overall_risk": "HIGH",
        "summary": {"vulnerability_count": 2, "confirmed_count": 1,
                    "by_severity": {"Critical": 1, "High": 1, "Medium": 0, "Low": 0}},
        "evidence_levels": {"level3_proven": 1, "level2_evidence": 2, "level1_observed": 3,
                            "level0_info": 5, "priority_actions": 3},
        "business_impact_summary": {"critical": 1, "high": 1, "medium": 0, "low": 0},
        "attack_surface_plan": {"summary": {"total_surfaces": 20, "auth_after_surfaces": 5,
                                            "by_type": {"Object Reference": 5}}},
        "ai_payload_plan": {"summary": {"total_input_points": 31, "ai_selected_input_points": 12,
                                        "planned_executions": 9, "executed": 9}},
        "attack_path_summary": {"total_paths": 4, "confirmed_paths": 1},
        "path_priority_summary": {"prioritized_paths": 2},
        "attack_paths": [
            {"path_id": "AP-001", "title": "인증 후 객체 접근(IDOR) 경로",
             "path_confidence": "Confirmed Path", "risk_score": 22,
             "steps": ["Entry Point: 인터넷", "Login Page: 로그인", "Object Reference: accountId",
                       "Evidence: Level 3", "Report Impact: 정보 노출"],
             "possible_impact": "타 사용자 정보 노출 가능"},
        ],
        "prioritized_paths": [{"path_id": "AP-001", "score": 22, "priority": "Critical"}],
        "priority_action_plan": [
            {"rank": "P1", "title": "객체 단위 접근 제어 검증 추가", "priority": "Critical",
             "team": "Backend", "effort": "1~3일", "benefit": "교차 계정 접근 차단"},
        ],
        "business_impact": [{"title": "관리자 IDOR 실증 — 객체 참조", "family": "idor",
                             "business_risk": "고객 정보 노출 가능",
                             "impact_category": ["Access Control", "Data Exposure"],
                             "potential_consequence": ["고객 정보 노출 가능"], "regulatory_risk": []}],
        "remediation_items": [{"finding_title": "관리자 IDOR 실증 — 객체 참조", "family": "idor",
                               "remediation_title": "객체 단위 접근 제어 검증 추가",
                               "estimated_effort": "1~3일", "responsible_team": "Backend",
                               "expected_benefit": "교차 계정 접근 차단"}],
        # 엔진 내부(부록으로 가야 함)
        "solver_summary": {"solver_runs": 3, "by_solver": {"idor_solver": 1}},
        "agent_results": [{"path_id": "AP-001", "family": "idor", "agents": [
            {"agent_name": "object_reference_agent", "agent_role": "객체 참조 분석",
             "observation": "객체 식별자 존재", "limitation": "읽기 전용", "recommended_next_step": "검증"}]}],
        "agent_summary": {"agent_runs": 3},
        "evidence_graph_summary": {"total_nodes": 40, "total_edges": 55, "evidence_nodes": 8,
                                   "agent_observations": 3},
        "good_items": [{"title": "TLS 1.2 적용"}],
        "asset_inventory": [{"ip": "10.0.0.5", "asset_type": "Database Server",
                             "open_ports": [3306], "detected_services": [{"family": "mysql"}],
                             "risk_tags": ["exposed_database"], "priority": "High", "os_guess": "linux",
                             "web_endpoints": []}],
        "asset_summary": {"total_assets": 1, "high_priority": 1, "internet_exposed": 0},
        "network_exposure_summary": {"input_targets": ["10.0.0.0/24"], "target_type": "cidr",
                                     "live_hosts": 1, "open_ports": 1, "services": 1,
                                     "web_candidates": 0, "high_priority_assets": 1},
        "findings": [{"title": "관리자 IDOR 실증 — 객체 참조", "confidence": "CONFIRMED_RESPONSE",
                      "severity": "HIGH", "evidence_detail": "타 계정 노출"}],
    }


# ── Security Score ───────────────────────────────────────────────────────────
def test_security_score_deterministic():
    a = _rich()
    s = report._v3_security_score(a)
    assert 0 <= s <= 100
    # 동일 입력 → 동일 점수
    assert report._v3_security_score(a) == s


# ── 통합 대시보드(6 그룹) ─────────────────────────────────────────────────────
def test_v3_dashboard_six_groups():
    doc = Document(); a = _rich()
    report._add_v3_dashboard(doc, a)
    txt = _txt(doc)
    assert "보안 진단 대시보드" in txt and "Security Score" in txt
    # v4: 한글 그룹명
    assert "위험도" in txt and "검증 결과" in txt and "자산" in txt and "네트워크" in txt


# ── Attack Path Flow (가로 흐름, 엔진 내부 없음) ─────────────────────────────
def test_v3_attack_path_flow():
    doc = Document(); a = _rich()
    report._add_v3_attack_path_flow(doc, a)
    txt = _txt(doc)
    assert "공격 경로" in txt and "경로 #1" in txt
    assert "➔" in txt                    # 가로 플로우 화살표
    assert "accountId" in txt
    # 엔진 내부(Solver/Agent) 미표시
    assert "Solver" not in txt and "Agent" not in txt


# ── 우선 조치 카드 ───────────────────────────────────────────────────────────
def test_v3_priority_action_cards():
    doc = Document(); a = _rich()
    report._add_v3_priority_actions(doc, a)
    txt = _txt(doc)
    assert "우선 조치 계획" in txt
    assert "P1" in txt and "객체 단위 접근 제어" in txt
    assert "Backend" in txt and "1~3일" in txt


# ── Developer Appendix: 엔진 내부는 부록에만 ─────────────────────────────────

def test_dev_appendix_can_be_disabled(monkeypatch):
    monkeypatch.setenv("REPORT_DEV_APPENDIX", "false")
    doc = Document(); a = _rich()
    report._add_developer_appendix(doc, a)
    assert all(not p.text for p in doc.paragraphs)


# ── 취약점 카드: 엔진 내부(Solver/Agent/Chain) 행 제거, 핵심 필드 유지 ───────
def test_vuln_card_no_engine_internals():
    a = _rich()
    f = a["findings"][0]
    doc = Document()
    report._add_observed_view_block(doc, f, a)
    txt = _txt(doc)
    assert "관찰 내용" in txt and "확보 증거" in txt
    assert "비즈니스 영향" in txt and "권장 조치" in txt
    assert "관련 Solver" not in txt and "관련 Agent" not in txt and "Evidence Chain" not in txt


# ── 전체 보고서 생성(데이터 유지) ────────────────────────────────────────────
def test_full_report_v3_intact(monkeypatch):
    monkeypatch.setenv("REPORT_SHOW_ENGINE_INTERNAL", "true")   # 풀 리포트(엔진내부 포함) 검증
    a = _rich()
    buf = report.generate_report({"domain": "t.example.com", "analysis": a,
                                  "created_at": "2026-06-30", "results": []})
    doc = Document(io.BytesIO(buf.getvalue()))
    txt = _txt(doc)
    assert "Security Assessment Report" in txt
    assert "점검 목적" in txt                           # v8 8단: 1.목적
    assert "취약점 상세" in txt                         # 6.상세
    assert "대응방안" in txt                            # 7.대응
    assert "보안 진단 대시보드" in txt                 # (부록으로 이동)
    assert "개선 로드맵" in txt                         # Roadmap
    assert "Developer Appendix" in txt                # 엔진 내부 분리
    assert "Security Score" in txt
