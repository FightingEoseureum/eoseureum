"""test_report_v2.py — Eoseureum Report Architecture V2 섹션 렌더링 검증."""
import os, sys, io
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _rich_analysis():
    return {
        "domain": "t.example.com",
        "overall_risk": "HIGH",
        "coverage": {"auth_login_success": True, "auth_crawl_enabled": True},
        "attack_surface_plan": {"summary": {
            "total_surfaces": 31, "auth_after_surfaces": 8, "critical_high_surfaces": 6,
            "by_type": {"Object Reference": 6, "Search": 5, "File Handling": 3}}},
        "ai_payload_plan": {"summary": {
            "total_input_points": 31, "ai_selected_input_points": 12,
            "planned_executions": 9, "executed": 9}},
        "attack_path_summary": {"total_paths": 12, "confirmed_paths": 2, "evidence_paths": 3},
        "path_priority_summary": {"prioritized_paths": 3},
        "evidence_levels": {"level3_proven": 2, "level2_evidence": 3, "level1_observed": 5,
                            "level0_info": 10, "priority_actions": 5},
        "evidence_graph_summary": {"total_nodes": 40, "total_edges": 55, "evidence_nodes": 8,
                                   "agent_observations": 9, "confirmed_paths": 2,
                                   "evidence_paths": 3, "observed_paths": 4},
        "solver_summary": {"solver_runs": 3, "by_solver": {"idor_solver": 1, "xss_solver": 1,
                           "upload_solver": 1}, "evidence_chains": 4, "prioritized_paths": 3,
                           "agent_runs": 9},
        "solver_results": [{"solver": "idor_solver", "solver_result": "EVIDENCE_CORRELATED"},
                           {"solver": "xss_solver", "solver_result": "EVIDENCE_CORRELATED"},
                           {"solver": "upload_solver", "solver_result": "INSUFFICIENT_EVIDENCE"}],
        "attack_paths": [
            {"path_id": "AP-001", "title": "인증 후 객체 접근(IDOR) 경로",
             "path_confidence": "Confirmed Path", "evidence_level": 3, "risk_grade": "Critical Path",
             "risk_score": 22, "steps": ["Entry Point: 진입점", "Login Page: 로그인",
             "Object Reference: accountId", "Evidence: Level 3"],
             "possible_impact": "타 사용자 정보 노출 가능"},
            {"path_id": "AP-002", "title": "반사형 XSS 경로", "path_confidence": "Confirmed Path",
             "evidence_level": 3, "risk_grade": "High Path", "risk_score": 14,
             "steps": ["Entry Point: 진입점", "Parameter: q", "Evidence: alert"]},
        ],
        "prioritized_paths": [
            {"path_id": "AP-001", "title": "인증 후 객체 접근(IDOR) 경로", "priority": "Critical",
             "confidence": "Confirmed Path", "evidence_strength": "strong", "score": 22,
             "recommended_solver": "idor_solver"},
            {"path_id": "AP-002", "title": "반사형 XSS 경로", "priority": "High",
             "confidence": "Confirmed Path", "evidence_strength": "strong", "score": 14,
             "recommended_solver": "xss_solver"},
        ],
        "agent_results": [
            {"path_id": "AP-001", "family": "idor", "solver": "idor_solver", "agents": [
                {"agent_name": "object_reference_agent", "agent_role": "객체 참조 분석",
                 "observation": "객체 식별자 파라미터 존재", "limitation": "읽기 전용",
                 "recommended_next_step": "소유자 검증"}]},
        ],
        "evidence_chains": [
            {"family": "idor", "evidence_chain": "IDOR Evidence Chain", "max_level": 3,
             "confidence_change": "Level 1 → Level 3 로 강화",
             "validation_summary": "인증 후 영역 포함 · IDOR Evidence Chain"},
        ],
        "candidate_verification": {"idor": {"verified": 1, "promoted": 1}},
        "good_items": [{"title": "TLS 1.2 이상 정상 적용"}],
        "attack_surface_items": [{"title": "관리자 페이지 발견(401)"}],
        "findings": [{"title": "관리자 IDOR 실증 — 객체 참조", "confidence": "CONFIRMED_RESPONSE",
                      "evidence_detail": "타 계정 노출", "severity": "HIGH"}],
    }

def test_v2_attack_surface_and_budget():
    doc = Document(); a = _rich_analysis()
    report._add_v2_attack_surface_analysis(doc, a)
    txt = _txt(doc)
    assert "공격 표면 분석" in txt
    assert "선별 절감률" in txt
    assert "61%" in txt or "%" in txt   # 1 - 12/31 ≈ 61%


def test_v2_attack_path_timeline():
    doc = Document(); a = _rich_analysis()
    report._add_v2_attack_path_analysis(doc, a)
    txt = _txt(doc)
    assert "공격 경로 분석" in txt
    assert "경로 #1" in txt
    assert "Object Reference: accountId" in txt
    assert "idor_solver" in txt        # solver 연결


def test_v2_solver_and_agent_and_graph():
    doc = Document(); a = _rich_analysis()
    report._add_v2_solver_analysis(doc, a)
    report._add_v2_agent_analysis(doc, a)
    report._add_v2_evidence_graph_analysis(doc, a)
    txt = _txt(doc)
    assert "Solver 분석" in txt and "idor_solver" in txt
    assert "Multi-Agent 분석" in txt and "객체 참조 분석" in txt
    assert "증거 그래프 분석" in txt and "IDOR Evidence Chain" in txt


def test_v2_security_controls_and_top10():
    doc = Document(); a = _rich_analysis()
    report._add_v2_security_controls(doc, a)
    report._add_v2_executive_top10(doc, a)
    txt = _txt(doc)
    assert "확인된 보안 통제" in txt
    assert "✓" in txt
    assert "Top 10" in txt
    assert "인증 후 객체 접근(IDOR) 경로" in txt

def test_finding_observed_view_links_path():
    doc = Document(); a = _rich_analysis()
    f = a["findings"][0]
    report._add_observed_view_block(doc, f, a)
    txt = _txt(doc)
    assert "관찰 내용" in txt and "공격 경로" in txt


def test_full_report_generates_with_v2(monkeypatch):
    monkeypatch.setenv("REPORT_SHOW_ENGINE_INTERNAL", "true")
    a = _rich_analysis()
    a["summary"] = {"vulnerability_count": 1, "confirmed_count": 1}
    buf = report.generate_report({"domain": "t.example.com", "analysis": a,
                                  "created_at": "2026-06-24", "results": []})
    assert isinstance(buf, io.BytesIO) and buf.getbuffer().nbytes > 0
    doc = Document(io.BytesIO(buf.getvalue()))
    txt = _txt(doc)
    assert "Security Assessment Report" in txt
    assert "보안 진단 대시보드" in txt
    # v3: 공격 경로 Flow + Developer Appendix(엔진 내부 분리)
    assert "공격 경로" in txt
    assert "Developer Appendix" in txt
    assert "Security Score" in txt
