"""test_service_pipeline_v2.py — Service Security Framework v2 (전체 파이프라인 통합)."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import service_pipeline as sp
import service_surface_planner as ssp
from agents import agent_registry as areg
import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _hosts():
    return [{"host": "t.example.com", "services": [
        {"port": 22, "service": "ssh", "banner": "SSH-2.0-OpenSSH_8.2 password"},
        {"port": 445, "service": "microsoft-ds", "banner": "SMB guest allowed"},
        {"port": 3306, "service": "mysql", "version": "5.7.30"},
        {"port": 6379, "service": "redis", "banner": "redis noauth protected mode off"},
        {"port": 2375, "service": "docker"},
    ]}]


# ── 1. Service Path Prioritization ───────────────────────────────────────────
def test_service_path_prioritization():
    plan = ssp.plan_service_surfaces(_hosts())
    prio = sp.prioritize_service_paths(plan["service_surfaces"], plan["service_attack_paths"])
    sel = prio["selected"]
    assert len(sel) == 5
    # 점수 내림차순
    assert [s["score"] for s in sel] == sorted([s["score"] for s in sel], reverse=True)
    # redis(무인증) / docker(컨테이너)가 High 이상
    redis = next(s for s in sel if s["family"] == "redis")
    assert redis["priority"] in ("Critical", "High")
    assert redis["recommended_solver"] == "redis_solver"
    assert prio["summary"]["high_priority"] >= 1


# ── 2. Service Solver Orchestration ──────────────────────────────────────────
def test_service_solver_orchestration():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    assert out["ran"] is True
    sres = out["service_solver_results"]
    assert len(sres) >= 1
    solvers = {r["solver"] for r in sres}
    assert "redis_solver" in solvers or "mysql_solver" in solvers
    # Solver 안전 불변식
    for r in sres:
        assert r["level_changed"] is False and r["did_discovery"] is False


# ── 3. Service Agent Group 실행 ──────────────────────────────────────────────
def test_service_agent_groups_run():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    ares = out["service_agent_results"]
    fams = {r["family"] for r in ares}
    assert "ssh" in fams or "redis" in fams
    ssh_res = next((r for r in ares if r["family"] == "ssh"), None)
    if ssh_res:
        names = {a["agent_name"] for a in ssh_res["agents"]}
        assert "remote_access_agent" in names
        for a in ssh_res["agents"]:
            assert a["did_discovery"] is False and a["level_changed"] is False


def test_service_agents_registered():
    for fam in ("ssh", "ftp", "smb", "mysql", "redis", "docker", "k8s", "smtp"):
        assert areg.select_agents(fam), f"{fam} agents missing"
    assert areg.select_agents("docker")[0].name == "docker_api_exposure_agent"


# ── 5. Service Evidence Graph ────────────────────────────────────────────────
def test_service_evidence_graph_nodes_edges():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    seg = out["service_evidence_graph"]
    types = {n["node_type"] for n in seg["nodes"]}
    assert sp.N_SURFACE in types and sp.N_PORT in types
    rels = {e["relation"] for e in seg["edges"]}
    assert sp.E_EXPOSED in rels
    assert any(r in rels for r in (sp.E_BY_AGENT, sp.E_TO_IMPACT, sp.E_TO_REMED))
    assert seg["summary"]["total_nodes"] > 0


# ── 6. Service Validation Readiness ──────────────────────────────────────────
def test_service_validation_candidates():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    cands = out["service_validation_candidates"]
    assert cands
    fams = {c["family"] for c in cands}
    # ssh(L1) 후보 포함, redis(L2)는 이미 증거 → 제외
    assert "ssh" in fams
    for c in cands:
        assert c["promotable_to_level3"] is False   # 서비스는 자동 exploit 없이 L2까지


# ── 7~8. Business Impact / Remediation ───────────────────────────────────────
def test_service_business_impact_and_remediation():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    imp = out["service_business_impact"]
    rem = out["service_remediation_items"]
    assert imp and rem
    fams_imp = {i["family"] for i in imp}
    assert "mysql" in fams_imp
    # 서비스 remediation KB 매핑
    mysql_rem = next((r for r in rem if r["family"] == "mysql"), None)
    assert mysql_rem and "DBA" in mysql_rem["responsible_team"]


# ── 10. 저장 구조 / 9. Controls ──────────────────────────────────────────────
def test_pipeline_output_keys():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    for k in ("service_prioritized_paths", "service_solver_results", "service_agent_results",
              "service_evidence_graph", "service_validation_candidates",
              "service_business_impact", "service_remediation_plan",
              "service_security_controls", "service_summary"):
        assert k in out


# ── Report 렌더링 ─────────────────────────────────────────────────────────────
def test_report_service_sections_render():
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    a = {"domain": "t",
         "service_surfaces": out["service_surfaces"],
         "service_surface_summary": out["service_summary"]["surface_summary"],
         "service_summary": out["service_summary"],
         "service_solver_results": out["service_solver_results"],
         "service_agent_results": out["service_agent_results"],
         "service_evidence_graph": out["service_evidence_graph"],
         "service_validation_candidates": out["service_validation_candidates"],
         "service_validation_summary": out["service_validation_summary"],
         "service_business_impact": out["service_business_impact"],
         "service_remediation_plan": out["service_remediation_plan"]}
    doc = Document()
    report._add_v2_dashboard(doc, a, "t")
    report._add_v2_service_solver_agent(doc, a)
    report._add_v2_service_evidence_graph(doc, a)
    report._add_v2_service_validation(doc, a)
    report._add_v2_service_remediation(doc, a)
    txt = _txt(doc)
    assert "서비스 Solver 분석" in txt
    assert "서비스 Multi-Agent 분석" in txt
    assert "서비스 증거 그래프" in txt
    assert "서비스 추가 검증 가능 항목" in txt
    assert "서비스 조치 계획" in txt
    assert "Svc Solver" in txt   # dashboard service KPI


def test_no_new_discovery_invariant():
    # 서비스 Solver/Agent 결과 전체에서 discovery/level 변경이 없어야 함
    out = sp.run_service_pipeline({"domain": "t"}, _hosts())
    assert out["service_summary"]["service_solver_runs"] >= 1
    for r in out["service_solver_results"]:
        assert r["did_discovery"] is False and r["level_changed"] is False
    for ar in out["service_agent_results"]:
        for a in ar["agents"]:
            assert a["did_discovery"] is False and a["level_changed"] is False
