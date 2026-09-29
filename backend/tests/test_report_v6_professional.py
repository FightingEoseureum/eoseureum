"""test_report_v6_professional.py — Professional Report Design v6."""
import os, sys, io
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report
import report_pdf
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _rich():
    from tests.test_report_v4_ux import _rich as r
    return r()


def _gen(analysis=None):
    a = analysis or _rich()
    buf = report.generate_report({"domain": "t.example.com", "analysis": a,
                                  "created_at": "2026-07-01", "results": []})
    return buf, Document(io.BytesIO(buf.getvalue()))


# ── 보고서 폰트 + fallback ────────────────────────────────────────────────────
def test_report_font_applied():
    buf, _ = _gen()
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    xml = z.read("word/document.xml").decode("utf-8", "ignore")
    assert report.FONT in xml
    # 서버에 지정 폰트가 있으면 그 폰트, 없으면 fallback 이름이라도 유효
    assert report.FONT


def test_font_fallback_when_missing(monkeypatch):
    monkeypatch.setenv("REPORT_FONT", "NoSuchFont_ZZZ")
    # fc-list 매칭 실패 → fallback 목록/우선값 반환(예외 없이)
    f = report._resolve_report_font()
    assert isinstance(f, str) and f


# ── Cover Layout / Brand ──────────────────────────────────────────────────────
def test_cover_v6_layout():
    _, doc = _gen(); txt = _txt(doc)
    assert "Security Assessment Report" in txt
    assert "Web · Service · Network Security Assessment" in txt
    assert "C O N F I D E N T I A L" in txt
    assert "Prepared by" in txt and "Eoseureum" in txt
    assert "Report No" in txt and "Version" in txt
    assert "종합 위험도" in txt   # v8: 위험도 배지(HTML/PDF 표지와 통일)
    assert "공격 표면" in txt      # v8: 심각도 KPI 스트립


# ── Executive Summary: 6항목 + Assessment Scope ─────────────────────────────
def test_executive_summary_v6():
    _, doc = _gen(); txt = _txt(doc)
    assert "보안 점수" in txt and "종합 위험도" in txt
    assert "Critical" in txt and "High" in txt and "비즈니스 영향" in txt
    assert "즉시 조치" in txt and "실증 확인" in txt
    assert "Assessment Scope" in txt
    assert "Scan Target" in txt and "Scan Period" in txt


# ── Dashboard 인포그래픽(분포 막대) ──────────────────────────────────────────
def test_dashboard_distribution_bars():
    doc = Document(); a = _rich()
    report._add_v3_dashboard(doc, a)
    txt = _txt(doc)
    assert "Severity Distribution" in txt
    assert "Validation Distribution" in txt
    assert "█" in txt or "░" in txt      # 막대 게이지


# ── Attack Story: Kill Chain 단계 ────────────────────────────────────────────

# ── Finding Card: 요약 + 4박스 + 참조 ────────────────────────────────────────
def test_finding_card_v6():
    a = _rich(); doc = Document()
    report._add_observed_view_block(doc, a["findings"][0], a)
    txt = _txt(doc)
    assert "요약:" in txt
    for box in ("관찰 내용", "확보 증거", "비즈니스 영향", "권장 조치"):
        assert box in txt
    assert "담당" in txt and "예상 공수" in txt and "참조" in txt


# ── PDF 지원(도구 없으면 graceful) ───────────────────────────────────────────
def test_pdf_convert_graceful():
    # LibreOffice 미설치 환경: None 반환(예외 없이) — 호출부는 DOCX 폴백
    out = report_pdf.convert_docx_to_pdf(b"not-a-real-docx")
    assert out is None or isinstance(out, bytes)
    assert isinstance(report_pdf.pdf_available(), bool)


def test_pdf_find_soffice_returns_path_or_none():
    p = report_pdf._find_soffice()
    assert p is None or isinstance(p, str)


# ── Developer Appendix Toggle ────────────────────────────────────────────────
_APPENDIX_HEADING = "Developer Appendix (엔진 내부 정보)"


def test_dev_appendix_disabled_not_generated(monkeypatch):
    monkeypatch.setenv("REPORT_DEV_APPENDIX", "false")
    _, doc = _gen()
    txt = _txt(doc)
    assert _APPENDIX_HEADING not in txt   # 실제 부록 섹션이 생성되지 않음


def test_dev_appendix_enabled_present(monkeypatch):
    monkeypatch.setenv("REPORT_SHOW_ENGINE_INTERNAL", "true")   # 개발자 부록은 엔진내부 옵션
    _, doc = _gen()
    txt = _txt(doc)
    assert _APPENDIX_HEADING in txt


# ── 전체 회귀(디자인 유지) ───────────────────────────────────────────────────
def test_full_report_v6_intact():
    _, doc = _gen(); txt = _txt(doc)
    for sect in ("Security Assessment Report", "Contents", "Executive Summary",
                 "보안 진단 대시보드", "점검 목적", "취약점 상세", "발견된 취약점",
                 "개선 로드맵", "부록"):
        assert sect in txt, f"섹션 누락: {sect}"


def test_engine_internal_hidden_by_default(monkeypatch):
    """A안: 엔진 내부/개발자 부록은 기본 숨김(간결 보고서). 토글 시에만 노출."""
    monkeypatch.delenv("REPORT_SHOW_ENGINE_INTERNAL", raising=False)
    _, doc = _gen(); txt = _txt(doc)
    assert _APPENDIX_HEADING not in txt          # Developer Appendix 기본 숨김
    assert "보안 진단 대시보드" in txt            # KPI 대시보드는 유지(핵심)
    # 토글 켜면 다시 노출
    monkeypatch.setenv("REPORT_SHOW_ENGINE_INTERNAL", "true")
    _, doc2 = _gen(); txt2 = _txt(doc2)
    assert _APPENDIX_HEADING in txt2
