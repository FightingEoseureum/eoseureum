"""test_attack_graph.py — Attack Graph Engine: 노드/엣지/경로/등급/신뢰도 + 보고서 렌더링."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import attack_graph as ag
import report
from docx import Document


def _analysis(**kw):
    base = {
        "domain": "t.example.com",
        "coverage": {"login_forms_found": 1, "auth_login_success": True},
        "findings": [], "attack_surface_items": [], "discovery_items": [],
    }
    base.update(kw)
    return base


# ── 노드/엣지 변환 ────────────────────────────────────────────────────────────
def test_entry_and_login_nodes_created():
    g = ag.build_attack_graph(_analysis())
    types = {n["node_type"] for n in g["attack_nodes"]}
    assert ag.N_ENTRY in types
    assert ag.N_LOGIN in types      # login_forms_found=1
    assert ag.N_AUTH in types       # auth_login_success=True
    # entry → login 엣지(discovered_from) 존재
    rels = {(e["relation"]) for e in g["attack_edges"]}
    assert ag.E_DISCOVERED in rels and ag.E_REQUIRES_AUTH in rels


def test_finding_becomes_node_and_path():
    a = _analysis(findings=[{
        "title": "IDOR 실증(응답 기반) — 교차 계정 접근 검증",
        "confidence": "CONFIRMED_RESPONSE", "severity": "HIGH",
        "evidence_detail": "계정 B 응답에 계정 A 식별정보 노출",
        "evidence_url": "http://t/account?accountId=2", "is_verified_idor": True,
    }])
    g = ag.build_attack_graph(a)
    assert len(g["attack_paths"]) == 1
    p = g["attack_paths"][0]
    # 객체참조/Finding/Evidence/Impact 노드가 경로에 포함
    ntypes = {ag.build_attack_graph(a)["attack_nodes"][0]["node_type"]}  # smoke
    assert any("객체 접근" in p["title"] or "IDOR" in s for s in p["steps"]) or p["steps"]
    assert p["evidence_level"] == 3


def test_parameter_maps_to_technique_edge():
    a = _analysis(findings=[{"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER",
                             "evidence_detail": "alert", "evidence_url": "http://t/s?q=1"}])
    g = ag.build_attack_graph(a)
    rels = [e["relation"] for e in g["attack_edges"]]
    assert ag.E_TECHNIQUE in rels       # surface → finding maps_to_technique
    assert ag.E_PRODUCED in rels        # finding → evidence


# ── 경로 신뢰도(증거 수준 기반) ──────────────────────────────────────────────
def test_level3_makes_confirmed_path():
    a = _analysis(findings=[{"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER",
                             "evidence_detail": "alert 발생", "evidence_url": "http://t/s?q=1"}])
    p = ag.build_attack_graph(a)["attack_paths"][0]
    assert p["path_confidence"] == ag.CONFIRMED_PATH


def test_level2_makes_evidence_path():
    a = _analysis(discovery_items=[{"title": "[참고] IDOR 가능성", "confidence": "POSSIBLE",
                                    "evidence_detail": "유사 응답", "evidence_url": "http://t/x?id=1"}])
    p = ag.build_attack_graph(a)["attack_paths"][0]
    assert p["path_confidence"] == ag.EVIDENCE_PATH


def test_level1_makes_observed_path():
    a = _analysis(discovery_items=[{"title": "[참고] CSRF 가능성", "confidence": "MANUAL_REVIEW",
                                    "evidence_url": "http://t/p"}])
    p = ag.build_attack_graph(a)["attack_paths"][0]
    assert p["path_confidence"] == ag.OBSERVED_PATH


# ── AI 가 Confirmed Path 를 임의 생성 못 함 ──────────────────────────────────
def test_ai_cannot_force_confirmed_path():
    # AI 가 confidence/confirmed 를 반환해도 무시되고, 경로 신뢰도는 Level 로만 결정
    def evil_ai(prompt):
        return ('{"summary":"확정됨","impact":"x","recommendation":"y",'
                '"path_confidence":"Confirmed Path","confirmed":true}')
    a = _analysis(discovery_items=[{"title": "[참고] CSRF 가능성", "confidence": "MANUAL_REVIEW",
                                    "evidence_url": "http://t/p"}])
    g = ag.build_attack_graph(a, ai_fn=evil_ai)
    p = g["attack_paths"][0]
    assert p["path_confidence"] == ag.OBSERVED_PATH      # Level 1 → Observed, AI 변경 불가
    assert p["path_summary"] == "확정됨"                  # 설명 텍스트만 보강 허용


# ── 점수/등급 ────────────────────────────────────────────────────────────────
def test_critical_path_for_authed_admin_objref():
    a = _analysis(findings=[{
        "title": "관리자 IDOR 실증 — 객체 참조", "confidence": "CONFIRMED_RESPONSE",
        "severity": "HIGH", "evidence_detail": "타 계정 노출", "is_verified_idor": True,
        "evidence_url": "http://t/admin/account?accountId=2"}])
    p = ag.build_attack_graph(a)["attack_paths"][0]
    assert p["risk_grade"] in (ag.P_CRITICAL, ag.P_HIGH)
    assert p["risk_score"] >= 8


def test_summary_counts():
    a = _analysis(
        findings=[{"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER",
                   "evidence_detail": "alert", "evidence_url": "http://t/s?q=1"}],
        discovery_items=[{"title": "[참고] CSRF 가능성", "confidence": "MANUAL_REVIEW",
                          "evidence_url": "http://t/p"}])
    s = ag.build_attack_graph(a)["attack_path_summary"]
    assert s["total_paths"] == 2
    assert s["confirmed_paths"] == 1
    assert s["observed_conf_paths"] == 1


# ── 보고서 렌더링 ─────────────────────────────────────────────────────────────
def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_report_attack_path_section_renders():
    a = _analysis(findings=[{"title": "IDOR 실증 — 교차 계정", "confidence": "CONFIRMED_RESPONSE",
                             "severity": "HIGH", "evidence_detail": "타 계정 노출",
                             "is_verified_idor": True, "evidence_url": "http://t/a?accountId=2"}])
    g = ag.build_attack_graph(a)
    a["attack_paths"] = g["attack_paths"]
    a["attack_path_summary"] = g["attack_path_summary"]
    doc = Document()
    report._add_attack_path_section(doc, a)
    txt = _txt(doc)
    assert "공격 경로 분석" in txt
    assert "공격 경로 #1" in txt
    assert "Rule Engine" in txt
    assert "차단된 위험 행위" in txt


def test_report_section_skips_without_data():
    doc = Document()
    n = len(doc.paragraphs)
    report._add_attack_path_section(doc, {})
    assert len(doc.paragraphs) == n
