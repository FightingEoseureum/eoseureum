"""test_business_impact_remediation.py — Business Impact Engine + Remediation Planner + Dashboard."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import business_impact_engine as bie
import remediation_planner as rp
import remediation_kb as rkb
import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _analysis():
    return {
        "findings": [
            {"title": "관리자 IDOR 실증 — 객체 참조", "confidence": "CONFIRMED_RESPONSE",
             "evidence_detail": "타 계정 노출", "is_verified_idor": True},
            {"title": "반사형 XSS", "confidence": "CONFIRMED_BROWSER", "evidence_detail": "alert"},
        ],
        "discovery_items": [
            {"title": "[참고] 비즈니스 로직 후보", "confidence": "MANUAL_REVIEW"},
        ],
        "attack_paths": [
            {"path_id": "AP-001", "title": "인증 후 객체 접근(IDOR) 경로",
             "risk_grade": "Critical Path", "path_confidence": "Confirmed Path"},
        ],
    }


# ── Impact Mapping ───────────────────────────────────────────────────────────
def test_impact_mapping_idor():
    f = {"title": "관리자 IDOR 실증 — 객체 참조", "confidence": "CONFIRMED_RESPONSE",
         "is_verified_idor": True}
    imp = bie.assess_finding(f)
    assert imp["family"] == "idor"
    assert bie.ACCESS_CONTROL in imp["impact_category"]
    assert bie.DATA_EXPOSURE in imp["impact_category"]
    assert any("고객 정보" in c for c in imp["potential_consequence"])
    assert imp["priority"] in ("Critical", "High")


def test_impact_mapping_families():
    fams = {}
    for title, fam in [("Stored XSS 실증", "xss"), ("SQL 인젝션", "sqli"),
                       ("SSRF 후보", "ssrf"), ("비즈니스 로직 후보", "business_logic")]:
        imp = bie.assess_finding({"title": title, "confidence": "POSSIBLE"})
        fams[fam] = imp
    assert bie.BUSINESS_LOGIC in fams["business_logic"]["impact_category"]
    assert "내부" in fams["ssrf"]["affected_asset"]


# ── Regulatory Mapping ───────────────────────────────────────────────────────
def test_regulatory_mapping_possibility_only():
    imp = bie.assess_finding({"title": "SQL 인젝션", "confidence": "CONFIRMED_RESPONSE"})
    assert imp["regulatory_risk"]
    # '관련 가능성' 수준만 — 단정 표현 없음
    assert all("관련 가능성" in r for r in imp["regulatory_risk"])


# ── Executive Risk Score ─────────────────────────────────────────────────────
def test_executive_risk_score_levels():
    high = bie.executive_risk_score(evidence_level=3, path_priority="Critical Path",
                                    authenticated=True, admin=True, sensitive=True,
                                    categories=[bie.DATA_EXPOSURE])
    assert high["level"] == "Critical"
    low = bie.executive_risk_score(evidence_level=0, path_priority="", categories=[])
    assert low["level"] == "Low"


def test_build_business_impact_summary():
    out = bie.build_business_impact(_analysis())
    assert out["summary"]["total"] >= 2
    # 정렬: 점수 내림차순
    scores = [it["executive_risk_score"] for it in out["items"]]
    assert scores == sorted(scores, reverse=True)
    assert "by_category" in out["summary"]


# ── Remediation Planner + KB ─────────────────────────────────────────────────
def test_remediation_planner_maps_kb():
    impact = bie.build_business_impact(_analysis())
    plan = rp.plan_remediation(impact)
    assert plan["items"]
    idor_item = next((i for i in plan["items"] if i["family"] == "idor"), None)
    assert idor_item is not None
    assert idor_item["responsible_team"] == "Backend"
    assert "접근 제어" in idor_item["remediation_title"]
    # 우선 조치 계획(P1..) 생성, family 중복 병합
    plan_ranks = [p["rank"] for p in plan["priority_action_plan"]]
    assert plan_ranks[0] == "P1"
    fams = [p["family"] for p in plan["priority_action_plan"]]
    assert len(fams) == len(set(fams))


def test_remediation_kb_extensible():
    rkb.register_remediation("graphql", {"title": "GraphQL 권한 검증", "team": "Backend"})
    assert "graphql" in rkb.all_families()
    assert rkb.get_remediation("graphql")["title"] == "GraphQL 권한 검증"


# ── Dashboard / Priority Action Plan rendering ───────────────────────────────
def test_executive_risk_dashboard_renders():
    a = _analysis()
    impact = bie.build_business_impact(a)
    a["business_impact"] = impact["items"]
    a["business_impact_summary"] = impact["summary"]
    plan = rp.plan_remediation(impact)
    a["priority_action_plan"] = plan["priority_action_plan"]
    a["remediation_summary"] = plan["summary"]
    a["remediation_items"] = plan["items"]
    # 대시보드/플랜 렌더
    doc = Document()
    report._add_v2_executive_risk_dashboard(doc, a)
    report._add_v2_priority_action_plan(doc, a)
    txt = _txt(doc)
    assert "경영진 위험 대시보드" in txt
    assert "Critical Risk" in txt
    assert "우선 조치 계획" in txt
    assert "P1" in txt
    assert "Backend" in txt


def test_finding_block_shows_business_impact():
    a = _analysis()
    impact = bie.build_business_impact(a)
    a["business_impact"] = impact["items"]
    plan = rp.plan_remediation(impact)
    a["remediation_items"] = plan["items"]
    doc = Document()
    report._add_observed_view_block(doc, a["findings"][0], a)
    txt = _txt(doc)
    assert "비즈니스 영향" in txt
    assert "권장 조치" in txt
    assert "예상 공수" in txt


def test_dashboard_skips_without_data():
    doc = Document()
    report._add_v2_executive_risk_dashboard(doc, {})
    report._add_v2_priority_action_plan(doc, {})
    assert all(not p.text for p in doc.paragraphs)


# ── AI Failure Fallback (해당 모듈은 ai_fn 미사용 결정적 — 예외 없이 동작) ─────
def test_deterministic_no_ai():
    out = bie.build_business_impact(_analysis())
    plan = rp.plan_remediation(out)
    assert out["items"] and plan["priority_action_plan"]
