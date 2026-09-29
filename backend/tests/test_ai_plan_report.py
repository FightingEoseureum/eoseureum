"""test_ai_plan_report.py — 보고서에 AI 계획/검증 + 차단 사유가 표시되는지 검증."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report
from docx import Document


def _doc_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_ai_plan_section_renders_with_block_reasons():
    analysis = {
        "ai_payload_plan": {
            "summary": {
                "total_input_points": 12, "ai_selected_input_points": 12,
                "ai_payload_candidates": 40, "validator_passed": 25,
                "validator_blocked": 15, "planned_executions": 25,
                "executed": 10, "confirmed": 2, "possible": 3, "manual_review": 5,
                "critical_candidates": 4, "rce_proof_mode": False, "time_based_enabled": False,
                "blocked_by_verdict": {"BLOCKED_SHELL": 6, "BLOCKED_TIME_BASED": 4,
                                       "BLOCKED_DESTRUCTIVE": 5},
            },
            "critical": [
                {"vuln_type": "idor", "param": "id", "priority_level": "CRITICAL", "priority_score": 13},
            ],
        }
    }
    doc = Document()
    report._add_ai_payload_plan_section(doc, analysis)
    txt = _doc_text(doc)
    assert "AI 기반 공격 기법 계획 및 검증" in txt
    assert "Validator 차단 수" in txt
    # 차단 사유가 표시되어야 한다
    assert "Time-based SQLi 차단" in txt and "셸/웹쉘/코드실행 차단" in txt
    # AI 가 아닌 Rule Engine 판정임을 명시
    assert "Rule Engine" in txt and "판정 권한이 없" in txt
    # Critical 후보 표
    assert "CRITICAL" in txt


def test_ai_plan_section_skipped_when_absent():
    doc = Document()
    before = len(doc.paragraphs)
    report._add_ai_payload_plan_section(doc, {})   # 데이터 없으면 아무것도 안 함
    assert len(doc.paragraphs) == before
