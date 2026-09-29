"""test_service_security.py — Service Security Framework v1."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import service_surface_planner as ssp
import solver_registry as registry
import report
from docx import Document


def _txt(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _hosts():
    return [{"host": "t.example.com", "services": [
        {"port": 22, "service": "ssh", "banner": "SSH-2.0-OpenSSH_8.2 password"},
        {"port": 445, "service": "microsoft-ds", "banner": "SMB guest allowed"},
        {"port": 3306, "service": "mysql", "version": "5.7.30"},
        {"port": 6379, "service": "redis", "banner": "redis noauth protected mode off"},
        {"port": 2375, "service": "docker"},
        {"port": 25, "service": "smtp"},
        {"port": 80, "service": "http"},   # 웹 — 서비스 표면 제외
    ]}]


# ── Service Classification ───────────────────────────────────────────────────
def test_service_classification():
    h = _hosts()[0]["services"]
    by = {}
    for svc in h:
        c = ssp.classify_service(svc)
        if c:
            by[c["family"]] = c
    assert by["ssh"]["attack_surface_type"] == ssp.S_AUTH
    assert by["smb"]["attack_surface_type"] == ssp.S_FILE
    assert by["mysql"]["attack_surface_type"] == ssp.S_DB
    assert by["redis"]["attack_surface_type"] == ssp.S_DATA
    assert by["docker"]["attack_surface_type"] == ssp.S_CONTAINER
    assert by["smtp"]["attack_surface_type"] == ssp.S_MESSAGING


def test_http_excluded_from_service_surface():
    assert ssp.classify_service({"port": 80, "service": "http"}) is None
    assert ssp.classify_service({"port": 443, "service": "https"}) is None


def test_recommended_solver_routing():
    c = ssp.classify_service({"port": 22, "service": "ssh"})
    assert c["recommended_solver"] == "ssh_solver"
    c2 = ssp.classify_service({"port": 6379, "service": "redis"})
    assert c2["recommended_solver"] == "redis_solver"


# ── Evidence Level (무인증/노출 단서 → Level2) ───────────────────────────────
def test_evidence_level_unauth_is_level2():
    c = ssp.classify_service({"port": 6379, "service": "redis",
                              "banner": "redis noauth protected mode off"})
    assert c["evidence_level"] == 2     # noauth/protected mode off → 증거(Level2)
    c2 = ssp.classify_service({"port": 22, "service": "ssh", "banner": "OpenSSH"})
    assert c2["evidence_level"] == 1    # 단순 노출 → 관찰(Level1)


# ── Service Planner / Paths ──────────────────────────────────────────────────
def test_plan_service_surfaces_and_paths():
    out = ssp.plan_service_surfaces(_hosts())
    s = out["service_surface_summary"]
    assert s["total_service_surfaces"] == 6   # http 제외
    assert s["authentication_surfaces"] == 1
    assert s["database_surfaces"] == 1
    assert s["container_surfaces"] == 1
    # 서비스 공격 경로 생성(Internet→Service→조건→위험)
    paths = out["service_attack_paths"]
    assert len(paths) == 6
    redis_path = next(p for p in paths if p["service_family"] == "redis")
    assert "Data Exposure" in redis_path["title"]
    assert any("무인증" in st for st in redis_path["steps"])


# ── Service Solver Registry ──────────────────────────────────────────────────
def test_service_solvers_registered():
    for s in ("ssh_solver", "ftp_solver", "smtp_solver", "smb_solver",
              "mysql_solver", "redis_solver", "docker_solver", "k8s_solver"):
        assert s in registry.all_solvers()
    # family 라우팅
    assert registry.select_solver("ssh").name == "ssh_solver"
    assert registry.select_solver("database") or registry.select_solver("mysql")


def test_service_solver_no_discovery_no_level_change():
    sol = registry.get_solver("redis_solver")
    out = sol.solve({"evidence": {"level": 2}, "attack_path": {"title": "redis"}})
    assert out["did_discovery"] is False
    assert out["level_changed"] is False


# ── Business Impact / Controls ───────────────────────────────────────────────
def test_service_business_impact():
    out = ssp.plan_service_surfaces(_hosts())
    imp = ssp.service_business_impact(out["service_surfaces"])
    fams = {i["family"]: i for i in imp}
    assert "mysql" in fams
    assert "Data Exposure" in fams["mysql"]["impact_category"]
    assert any("민감 정보" in c for c in fams["mysql"]["potential_consequence"])
    # SSH → 관리자/원격 접근 위험
    assert "Access Control" in fams["ssh"]["impact_category"]


def test_service_controls_evidence_based():
    surfaces = [
        {"family": "smb", "attack_surface_type": ssp.S_FILE,
         "evidence": {"signing": "enabled"}},
        {"family": "ssh", "attack_surface_type": ssp.S_AUTH,
         "evidence": {"auth_hint": "publickey"}},
    ]
    controls = ssp.service_controls(surfaces)
    assert "SMB Signing 활성" in controls
    assert any("SSH Password Login 비활성" in c for c in controls)


# ── Dashboard / Report 표시 ───────────────────────────────────────────────────
def test_dashboard_and_report_render_service():
    out = ssp.plan_service_surfaces(_hosts())
    a = {"domain": "t.example.com",
         "service_surfaces": out["service_surfaces"],
         "service_surface_summary": out["service_surface_summary"],
         "service_controls": ["SMB Signing 활성"]}
    doc = Document()
    report._add_v2_dashboard(doc, a, "t.example.com")
    report._add_v2_service_security(doc, a)
    txt = _txt(doc)
    assert "Service Surface" in txt
    assert "서비스 보안 분석" in txt
    assert "Database" in txt
    # 보안 통제에 서비스 통제 포함
    controls = report._derive_verified_controls(a)
    assert "SMB Signing 활성" in controls
